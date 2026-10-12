"""Conservative inventory of thesis source archives; no raw market rows exported.

This tool DOES NOT reconstruct financial-filing availability or prove model
performance. It inspects candidate tabular schemas without extracting archives.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

MAX_MEMBER_PREVIEW = 256 * 1024
MAX_WORKBOOK_BYTES = 16 * 1024 * 1024
MAX_NESTED_ZIP_BYTES = 32 * 1024 * 1024
SKIP_SUFFIXES = {".pdf", ".docx", ".pptx", ".png", ".jpg", ".jpeg", ".gif", ".mp4", ".svg", ".html", ".txt"}
SUPPORTED = {".csv", ".tsv", ".xlsx", ".xlsm", ".json", ".parquet"}

IDENTIFIERS = {"company", "company_id", "ticker", "symbol", "security", "isin", "nse_symbol", "nse"}
ASOF = {"prediction_asof", "asof", "as_of_date", "signal_date", "prediction_date"}
TIME_KEYS = {"date", "datetime", "quarter", "quarter_end", "fiscal_quarter", "period", "year"}
LABELS = {"target_label", "direction", "target", "next_quarter_direction", "label", "next_quarter_up"}
PUBLICATION = {"published_at", "release_date", "filing_date", "publication_date", "available_at"}
OUTCOME_AVAILABLE = {"label_available_at", "outcome_available_at"}
FEATURE_NAMES = {
    "revenue", "total_revenue", "total_assets", "net_income", "eps", "roa", "roe",
    "debt_to_equity", "net_profit", "operating_income", "cash_flow", "market_cap",
    "volume", "close", "adj_close", "adjusted_close", "price", "volatility",
    "ebitda", "return", "momentum", "pe_ratio", "p_e", "leverage", "equity",
}
ARCHIVE = {".zip"}
VALID_SUFFIXES = SUPPORTED | ARCHIVE


def normalize(field: object) -> str:
    x = str(field or "").strip().lower()
    x = re.sub(r"[^a-z0-9]+", "_", x)
    return x.strip("_")


def discover_header(sample: bytes, suffix: str) -> list[str]:
    """Examine first 12 records, selecting the header with most relevant terms."""
    if sample.startswith(b"\xef\xbb\xbf"):
        sample = sample[3:]
    text = sample.decode("utf-8-sig", errors="replace").replace("\x00", "")
    sep = "\t" if suffix == ".tsv" else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=sep))[:12]
    candidates = []
    tokens = IDENTIFIERS | ASOF | TIME_KEYS | LABELS | PUBLICATION | OUTCOME_AVAILABLE | FEATURE_NAMES
    for row in rows:
        fields = [normalize(v) for v in row[:80]]
        signal = len(set(fields) & tokens)
        signal += int(len(fields) >= 3)
        if fields:
            candidates.append((signal, fields))
    if not candidates:
        return []
    return max(candidates, key=lambda x: x[0])[1]


def preview_workbook(data: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [{"sheet": None, "headers": [], "issue": "openpyxl_unavailable"}]
    output = []
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets[:12]:
            raw_rows = []
            for row in ws.iter_rows(min_row=1, max_row=16, max_col=80, values_only=True):
                raw_rows.append(list(row))
            ranked = []
            for row in raw_rows:
                h = [normalize(v) for v in row]
                signal = len(set(h) & (IDENTIFIERS | ASOF | TIME_KEYS | LABELS | PUBLICATION | OUTCOME_AVAILABLE | FEATURE_NAMES))
                if len([x for x in h if x]) > 2:
                    signal += 1
                ranked.append((signal, h))
            h = max(ranked, key=lambda x:x[0])[1] if ranked else []
            output.append({"sheet": str(ws.title)[:120], "headers": h})
        wb.close()
    except Exception as exc:
        return [{"sheet":None, "headers":[], "issue":"unreadable_workbook:"+type(exc).__name__}]
    return output


def grade(headers: list[str], path: str) -> dict[str, Any]:
    h = {normalize(x) for x in headers if x}
    id_col = bool(h & IDENTIFIERS)
    asof_col = bool(h & ASOF)
    time_col = bool(h & TIME_KEYS)
    label_col = bool(h & LABELS)
    pub_col = bool(h & PUBLICATION)
    label_time = bool(h & OUTCOME_AVAILABLE)
    predictor_count = len(h & FEATURE_NAMES)
    combined_time = bool(asof_col or time_col)
    score = 4*int(id_col) + 3*int(asof_col) + 2*int(time_col) + 5*int(label_col)
    score += 4*int(pub_col) + 2*int(label_time) + min(5, predictor_count)
    low = path.lower()
    for word in ("panel","thesis","dataset","merged","master","model_input","feature","target"):
        if word in low:
            score += 1
    if not (id_col and combined_time):
        classification = "raw_or_other_table"
    elif label_col and asof_col and pub_col and label_time:
        classification = "candidate_point_in_time_panel_schema_only"
    elif label_col:
        classification = "candidate_labelled_panel_missing_proof"
    else:
        classification = "candidate_company_time_table"
    issues = []
    if not id_col: issues.append("company_identifier_not_recognized")
    if not combined_time: issues.append("time_identifier_not_recognized")
    if not label_col: issues.append("outcome_label_not_recognized")
    if not asof_col: issues.append("prediction_asof_not_recognized")
    if not pub_col: issues.append("public_release_timestamp_not_recognized")
    if not label_time: issues.append("label_availability_timestamp_not_recognized")
    if label_col and not pub_col:
        issues.append("FISCAL_PERIOD_IS_NOT_FILING_PUBLICATION_DATE")
    issues.append("COLUMN_NAMES_DO_NOT_PROVE_VALID_RELEASE_TIMES_OR_CORPORATE_ACTIONS")
    return {
        "score": score, "classification": classification,
        "recognized": sorted(h & (IDENTIFIERS | ASOF | TIME_KEYS | LABELS | PUBLICATION | OUTCOME_AVAILABLE | FEATURE_NAMES)),
        "issues": issues,
    }


def inspect_file(zf: zipfile.ZipFile, member: zipfile.ZipInfo, virtual_path: str) -> list[dict[str, Any]]:
    suffix = Path(member.filename).suffix.lower()
    result = []
    base = {"path": virtual_path, "uncompressed_bytes": member.file_size, "kind": suffix}
    if suffix not in SUPPORTED:
        return result
    if suffix in (".csv", ".tsv"):
        with zf.open(member) as stream:
            sample = stream.read(MAX_MEMBER_PREVIEW)
        tabs = [{"sheet":None,"headers":discover_header(sample, suffix)}]
    elif suffix in (".xlsx", ".xlsm"):
        if member.file_size > MAX_WORKBOOK_BYTES:
            return [{**base, "status": "uninspected_large_workbook", "issue":"workbook_size_cap"}]
        with zf.open(member) as stream:
            raw = stream.read(MAX_WORKBOOK_BYTES + 1)
        if len(raw)>MAX_WORKBOOK_BYTES:
            return [{**base,"status":"uninspected_large_workbook","issue":"workbook_size_cap"}]
        tabs = preview_workbook(raw)
    elif suffix == ".json":
        with zf.open(member) as stream:
            sample = stream.read(MAX_MEMBER_PREVIEW)
        try:
            parsed = json.loads(sample)
            if isinstance(parsed,dict) and isinstance(parsed.get("data"),list) and parsed["data"]:
                parsed=parsed["data"]
            if isinstance(parsed,list) and parsed and isinstance(parsed[0],dict):
                h=list(parsed[0].keys())[:80]
            elif isinstance(parsed,dict):
                h=list(parsed.keys())[:80]
            else:
                h=[]
        except (ValueError,UnicodeDecodeError):
            h=[]
        tabs=[{"sheet":None,"headers":h}]
    else:
        return [{**base,"status":"schema_not_read","issue":"parquet_reader_not_available_without_opt_in"}]
    for tab in tabs:
        graded = grade(tab.get("headers") or [],virtual_path)
        result.append({
            **base,"sheet":tab.get("sheet"),"status":"header_inspected" if "issue" not in tab else "schema_unavailable",
            "headers":(tab.get("headers") or [])[:80], "schema_issue":tab.get("issue"),
            **graded
        })
    return result


def audit_archive(
    archive: Path,
    *,
    max_members: int = 25000,
    max_depth: int = 1,
) -> dict[str, Any]:
    """Read archives without unpacking; enforce size limits on nested ZIPs."""
    archive=Path(archive).expanduser()
    if not archive.is_file():
        raise FileNotFoundError(str(archive))
    if not zipfile.is_zipfile(archive):
        raise ValueError("Expected .zip archive")
    records: list[dict[str,Any]]=[]
    counts=Counter()
    warnings=[]
    seen=0

    def traverse(zf: zipfile.ZipFile, prefix: str, depth: int) -> None:
        nonlocal seen
        for info in zf.infolist():
            if info.is_dir():
                continue
            seen += 1
            counts["total_members"] += 1
            if seen > max_members:
                warnings.append("max_members_reached")
                return
            ext=Path(info.filename).suffix.lower()
            counts[ext or "<none>"] += 1
            virtual = f"{prefix}::{info.filename}" if prefix else info.filename
            if ext == ".zip" and depth < max_depth:
                if info.file_size > MAX_NESTED_ZIP_BYTES:
                    counts["nested_zips_skipped_large"] += 1
                    continue
                try:
                    with zf.open(info) as stream:
                        raw=stream.read(MAX_NESTED_ZIP_BYTES+1)
                    if len(raw)>MAX_NESTED_ZIP_BYTES:
                        counts["nested_zips_skipped_large"] += 1
                        continue
                    with zipfile.ZipFile(io.BytesIO(raw)) as inner:
                        traverse(inner,virtual,depth+1)
                except (OSError,ValueError,zipfile.BadZipFile,RuntimeError) as exc:
                    warnings.append(f"nested_zip_unreadable:{virtual}:{type(exc).__name__}")
                continue
            if ext not in SUPPORTED:
                continue
            try:
                records.extend(inspect_file(zf,info,virtual))
            except (OSError,ValueError,zipfile.BadZipFile,RuntimeError) as exc:
                records.append({
                    "path":virtual,"kind":ext,"status":"read_error",
                    "issue":type(exc).__name__
                })
        return

    with zipfile.ZipFile(archive) as zf:
        traverse(zf,archive.name,0)
    candidates=sorted(records, key=lambda r: (-r.get("score",0),r.get("path","")))[:125]
    classifications=dict(Counter(r.get("classification","unknown") for r in records))
    info={
        "generated_utc":datetime.now(timezone.utc).isoformat(),
        "source_archive_name":archive.name,
        "source_archive_size_bytes":archive.stat().st_size,
        "archives_inspected_without_extraction":True,
        "is_empirical_thesis_result":False,
        "is_source_verification":False,
        "total_members_visited":seen,
        "by_extension":dict(sorted(counts.items())),
        "records_inspected":len(records),
        "classifications":classifications,
        "priority_candidates":candidates,
        "warnings":warnings[:100],
        "critical_blockers":[
            "No sampled columns or filenames can establish public availability dates.",
            "Resolve STAR.NS corporate actions using independently checked raw price history.",
            "Confirm label definition, adjusted returns, split factors, and reporting date cutoffs.",
            "Prove each issuer's point-in-time filing availability, revisions and identifier mapping.",
            "Do not publish licensed S&P/Bloomberg data extracted from these ZIPs.",
            "Large nested ZIPs, large workbooks and Parquet may require targeted separate inspection."
        ],
    }
    return info


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive",type=Path,required=True)
    parser.add_argument("--out",type=Path,required=True)
    parser.add_argument("--max-members",type=int,default=25000)
    parser.add_argument("--max-depth",type=int,choices=(0,1,2),default=1)
    args=parser.parse_args(argv)
    if args.max_members<1:
        parser.error("--max-members must be positive")
    result=audit_archive(args.archive,max_members=args.max_members,max_depth=args.max_depth)
    args.out.parent.mkdir(parents=True,exist_ok=True)
    args.out.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps({
        "archive":result["source_archive_name"],
        "members":result["total_members_visited"],
        "tables_or_sheets":result["records_inspected"],
        "schema_classes":result["classifications"],
        "report":str(args.out)
    },indent=2))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
